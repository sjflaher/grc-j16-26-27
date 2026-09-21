#!/usr/bin/env python3
"""Build GRC J16 2026/27 intro decks (parents + boys) as .pptx for Google Slides import."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

MAROON = RGBColor(0x7A, 0x1F, 0x2B)   # Galway maroon
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF2, 0xEF, 0xEA)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def title_slide(prs, title, subtitle, notes=None):
    s = _blank(prs)
    _bg(s, MAROON)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = WHITE
    sub = s.shapes.add_textbox(Inches(0.8), Inches(4.1), Inches(11.7), Inches(1.2))
    tf2 = sub.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(24)
    p2.font.color.rgb = LIGHT
    _notes(s, notes)
    return s


def statement_slide(prs, text, kicker=None, notes=None):
    """Big single-statement slide."""
    s = _blank(prs)
    _bg(s, MAROON)
    if kicker:
        k = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.8))
        pk = k.text_frame.paragraphs[0]
        pk.text = kicker.upper()
        pk.font.size = Pt(20)
        pk.font.bold = True
        pk.font.color.rgb = LIGHT
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.7), Inches(2.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    _notes(s, notes)
    return s


def bullets_slide(prs, title, items, notes=None, title_size=32, body_size=20):
    """items: list of str or (str, level) or (str, level, bold)."""
    s = _blank(prs)
    _bg(s, WHITE)
    t = s.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.9))
    pt = t.text_frame.paragraphs[0]
    pt.text = title
    pt.font.size = Pt(title_size)
    pt.font.bold = True
    pt.font.color.rgb = MAROON
    body = s.shapes.add_textbox(Inches(0.9), Inches(1.55), Inches(11.6), Inches(5.5))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, str):
            text, level, bold = item, 0, False
        elif len(item) == 2:
            (text, level), bold = item, False
        else:
            text, level, bold = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = ("• " if level == 0 else "– ") + text if text else ""
        p.level = level
        p.font.size = Pt(body_size if level == 0 else body_size - 2)
        p.font.bold = bold
        p.font.color.rgb = DARK if level == 0 else GREY
        p.space_after = Pt(10)
    _notes(s, notes)
    return s


def table_slide(prs, title, headers, rows, notes=None, col_widths=None, body_size=16):
    s = _blank(prs)
    _bg(s, WHITE)
    t = s.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.9))
    pt = t.text_frame.paragraphs[0]
    pt.text = title
    pt.font.size = Pt(32)
    pt.font.bold = True
    pt.font.color.rgb = MAROON
    nrows, ncols = len(rows) + 1, len(headers)
    height = min(Inches(0.5) * nrows, Inches(5.4))
    shape = s.shapes.add_table(nrows, ncols, Inches(0.7), Inches(1.55), Inches(11.9), height)
    tbl = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = int(Inches(11.9) * w / total)
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = h
        for p in c.text_frame.paragraphs:
            p.font.size = Pt(body_size)
            p.font.bold = True
            p.font.color.rgb = WHITE
        c.fill.solid()
        c.fill.fore_color.rgb = MAROON
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i + 1, j)
            c.text = str(val)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(body_size)
                p.font.color.rgb = DARK
            c.fill.solid()
            c.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
    _notes(s, notes)
    return s


# ============================================================ PARENTS DECK
prs = new_deck()

title_slide(prs, "GRC Mens J16 — 2026/27 Season", "Parent Briefing — 30 August 2026")

bullets_slide(prs, "Agenda", [
    "Last season",
    "This season's goal",
    "How we'll get there: principles",
    "The training year: blocks & weekly schedule",
    "Race calendar & key dates",
    "What we need from parents",
    "What the boys own",
    "Tracking & technology",
])

bullets_slide(prs, "Last Season (25/26)", [
    ("Wins at Castleconnell and Athlone regattas", 0, True),
    "[ADD: 1–2 more results/highlights — IIRC, heads, crew development]",
    "A squad that finished the season faster, fitter and hungrier than it started",
    "Foundation laid: aerobic base, technical fundamentals, race experience",
], notes="Keep this short and warm. The message: last year worked, this year we build on it.")

statement_slide(prs, "Push for a Junior 16\nNational Championship.", kicker="This Season's Goal",
                notes="Last year: become nationally competitive at J15 — done. This year: convert competitive into winning. Everything in this programme is designed backwards from Champs, July 2027.")

bullets_slide(prs, "The Honest Starting Point — Fighting Headwinds", [
    "At GRC we face structural headwinds vs the biggest clubs:",
    ("Smaller athlete pool", 1),
    ("Boys spread across different schools, some out of town", 1),
    ("Limited scope for mid-week water sessions", 1),
    "",
    ("Our answer: we will not out-volume anyone — we will out-think, out-prepare and out-execute them.", 0, True),
    "Every session smart, efficient and intentional",
], notes="This is also why the parent logistics piece matters — the sessions we CAN do together have to happen and have to count.")

table_slide(prs, "The Step Up: J15 → J16",
    ["", "25/26 (J15)", "26/27 (J16)"],
    [
        ["Total volume", "~305 hrs", "~400 hrs"],
        ["Average per week", "~6 hrs", "~9 hrs"],
        ["Sessions per week", "5 (+ optionals)", "7–10"],
        ["New this year", "—", "Morning ergs, running, 2x S&C, sweep + sculling"],
    ],
    notes="400 hrs/year moves us from 'Club' toward 'National' on the training-volume ladder (McNeely: National = 600–800 hrs; Club/HS = 300–500). A meaningful commitment step — deliberately so, in a championship year. Reuse last year's McNeely table slide after this one if useful.",
    col_widths=[2, 2, 4])

table_slide(prs, "The Training Year: 5 Blocks",
    ["Block", "Dates", "Focus"],
    [
        ["Block 1", "Sep – Oct", "Re-entry, base building, baseline testing"],
        ["Block 2", "Late Oct – Dec", "Aerobic development, indoor season"],
        ["Block 3", "Jan – Mar", "Threshold work, head racing"],
        ["Block 4", "Mar – May", "Speed development, early regattas"],
        ["Block 5A/5B", "May – Jul", "Race prep → CHAMPIONSHIPS"],
    ],
    notes="Every block: detailed programme published at the start, 1:1 goal-setting meeting with each athlete, testing + review at the end.",
    col_widths=[2, 2.5, 5])

table_slide(prs, "A Typical Week (winter shape, Blocks 2–3)",
    ["Day", "Session"],
    [
        ["Monday", "Recovery / catch-up"],
        ["Tuesday", "Morning erg + evening S&C"],
        ["Wednesday", "Run"],
        ["Thursday", "Morning erg + evening S&C"],
        ["Friday", "Rest"],
        ["Saturday", "Water (sculling / sweep)"],
        ["Sunday", "Water x2 groups (sculling / sweep)"],
    ],
    notes="Morning sessions are the big logistical change this year. Exact schedule varies by block — full programme shared with all families.",
    col_widths=[2, 7], body_size=14)

bullets_slide(prs, "Race Calendar 2026/27", [
    ("Autumn / Winter", 0, True),
    ("Lough Rinn testing w/e — 19–20 Sep   •   St. Michael's HOR — 10 Oct", 1),
    ("Dublin HOR (tbc) — 14 Nov   •   PIRC — 21 Nov   •   IIRC — 16 Jan", 1),
    ("Spring Heads", 0, True),
    ("St. Michael's HOR — Feb   •   Galway HOR — 13–14 Mar", 1),
    ("Summer Regattas", 0, True),
    ("Neptune / Commercial — 3–4 Apr (Easter)   •   Limerick — 24 Apr", 1),
    ("Lough Rinn — 8–9 May   •   Metro — 22 May   •   Galway — 5 Jun", 1),
    ("Cork — 19 Jun   •   1K Classic / Lough Rinn — 26–27 Jun   •   Super Comp — 3–5 Jul", 1),
    ("CHAMPS — 9–11 July", 0, True),
], notes="Dates from provisional schedule — subject to confirmation.", body_size=18)

bullets_slide(prs, "What We Need From Parents", [
    ("Logistics — especially the new morning erg sessions (Tue/Thu) and race-day travel", 0, True),
    "The calendar — availability is locked in at the start of each block; please plan family commitments around published race dates where possible",
    "Support the standards at home — sleep, food, hydration (details with the boys)",
    "Let the boys own their rowing — next slide",
])

bullets_slide(prs, "What The Boys Own — Athlete-Led Standards", [
    "This year, responsibility sits with the athletes:",
    ("Attendance & communication: if a boy can't make a session, HE tells the coach, in good time — not a parent", 1),
    ("Nutrition, hydration, sleep, session prep and recovery", 1),
    ("Personal targets for every single session", 1),
    ("Block goals, set 1:1 with the coach", 1),
    "",
    ("Parents get the schedule and the race calendar. The boys get everything else.", 0, True),
], notes="Frame positively — this is athlete development, not parent exclusion. It's also what senior/college programmes will expect of them.")

bullets_slide(prs, "Tracking, Targets & Technology", [
    ("Individual targets replace crew averages this year", 0, True),
    "Every boy has personal erg, strength and technique targets per block",
    "Every session has a defined target for every athlete",
    "Data captured via Lumin app + heart-rate monitors (Polar)",
    "New this season: a squad web portal — boys log results, see personal targets and progression (in development)",
    "Close monitoring = early, supportive intervention where needed — progress is the motivator",
])

bullets_slide(prs, "Questions", [
    "Contact: [phone / email]",
    "Schedule & race calendar: [link to Sheet]",
])

prs.save("/Users/simonflaherty/Documents/GRC/01-season-2026-27/GRC J16 2026-27 - Parent Briefing.pptx")
print("parents deck saved")

# ============================================================ BOYS DECK
prs = new_deck()

title_slide(prs, "GRC Mens J16 — 2026/27", "This is a championship year.")

bullets_slide(prs, "Last Season", [
    ("Wins at Castleconnell and Athlone", 0, True),
    "[ADD 1–2 more highlights]",
    "",
    ("You proved you can compete. This year we find out if you can win.", 0, True),
])

statement_slide(prs, "A Junior 16\nNational Championship.", kicker="The Goal",
                notes="Let it sit. One line, no bullets.")

bullets_slide(prs, "The Truth About Where We Start", [
    "The clubs we have to beat:",
    ("Have more rowers than us", 1),
    ("Have everyone in one school", 1),
    ("Get more mid-week water time", 1),
    "",
    ("They can afford wasted sessions. We can't.", 0, True),
    ("And that discipline — never wasting a session — is exactly how we beat them.", 0, True),
], notes="Headwinds reframed as identity/edge.")

bullets_slide(prs, "Laying Bricks", [
    "Think of the season as building a wall. Every session = one brick.",
    ("~400 hours this season ≈ 306 bricks", 1),
    ("400 hours of training makes us competitive", 1),
    ("Winning means every brick is STONE, not cardboard", 1, True),
    "",
    "A stone brick =",
    ("1. Turn up PREPPED — rested, fed, hydrated", 1),
    ("2. FULL EFFORT — physical and mental, hit your target", 1),
    ("3. RECOVER properly — sleep and food are the mortar. A great session with no recovery is a brick with no mortar.", 1),
], notes="Origin story — grain of rice in a jar, and its flaw: the rice drops in whatever the quality. Bricks are different: quality is the whole point.")

bullets_slide(prs, "Three Questions After Every Session", [
    ("1.  Did I arrive prepped?", 0, True),
    ("2.  Did I hit my target?", 0, True),
    ("3.  Will I recover properly?", 0, True),
    "",
    ("3 yeses = a stone brick in the wall.", 0, True),
], body_size=26)

bullets_slide(prs, "The Season: 5 Blocks", [
    "Block 1 (Sep–Oct) → Block 2 (Nov–Dec) → Block 3 (Jan–Mar) → Block 4 (Mar–May) → Block 5 (May–CHAMPS)",
    "",
    "At the START of every block, each of you meets me 1:1:",
    ("Set your goals for the block", 1),
    ("Sort your schedule — then your availability is LOCKED IN", 1),
    ("Agree any personal tweaks to the programme", 1),
    "At the END of every block: testing + review — you'll see exactly what the block built",
    "",
    ("Missing a session you committed to should be extremely rare. Life happens — but committed means committed.", 0, True),
])

bullets_slide(prs, "Your Targets, Not Crew Averages", [
    ("No crew averages this year. Every one of you gets personal targets:", 0, True),
    ("Erg: wattage targets — per block, per session", 1),
    ("S&C: weight / rep progressions", 1),
    ("Technique: specific refinements to make each block", 1),
    "",
    "Every session has a target: a wattage to hold (erg), weights/reps to hit (S&C), 1 personal + 1 crew technique focus + boat speed (water)",
    "",
    ("You should never sit on an erg or get in a boat without knowing exactly what you're trying to do.", 0, True),
], notes="Targets derived from your own test results — they move when you improve. Progress against your own numbers is the scoreboard.")

bullets_slide(prs, "You Run This — Athlete-Led Standards", [
    "Your parents get two things: the schedule and the race calendar.",
    ("Everything else is yours:", 0, True),
    ("Can't make training? YOU message me, in good time. Not your mum. Not your dad.", 1),
    ("Your food, your water bottle, your sleep, your kit", 1),
    ("Your prep before sessions, your recovery after", 1),
    ("Knowing your targets before you walk in the door", 1),
], notes="Last year this was a goal; this year it's the standard. We can't afford passengers — including passengers in their own rowing.")

bullets_slide(prs, "The Week (winter shape)", [
    "Tue: Morning erg + evening S&C",
    "Wed: Run",
    "Thu: Morning erg + evening S&C",
    "Sat: Water",
    "Sun: Water",
    "",
    ("~9 hrs/week (last year ~6). Mornings are new — that's the step up to championship level.", 0, True),
    "National-level juniors train 600+ hrs/year; we're moving from ~305 to ~400 — and making every hour count double.",
])

bullets_slide(prs, "Racing", [
    "Testing w/e: Lough Rinn, 19–20 Sep",
    "St. Michael's HOR: Oct + Feb",
    "PIRC: 21 Nov  →  IIRC: 16 Jan",
    "Galway HOR: 13–14 Mar",
    "Regattas: Neptune/Commercial, Limerick, Lough Rinn, Metro, Galway, Cork, 1K Classic, Super Comp",
    ("CHAMPS: 9–11 July", 0, True),
    "",
    "Regular racing = regular proof of what the wall looks like.",
])

bullets_slide(prs, "Sleep, Food, Recovery — The Mortar", [
    "Protein ~2g per kg bodyweight; carbs AND fats matter",
    "Fruit & veg — \"Eat food. Not too much. Mostly plants.\" / 30 plants a week",
    "7–9 hours sleep, consistent times, wind-down routine",
    "HR monitor on for sessions — resting HR & HRV tell us when to push and when to back off",
])

bullets_slide(prs, "Tools", [
    "Lumin — session plans and logging",
    "Polar / HR monitor — every erg and water session",
    ("NEW: squad portal — log your results, see your targets, watch your wall get built brick by brick (coming this season)", 0, True),
])

statement_slide(prs, "One wall. 306 bricks.\nStarts Monday.", kicker="",
                notes="First brick: [first session details]. End on the analogy — optionally show an empty wall graphic that fills in over the season; same visual the portal will use.")

prs.save("/Users/simonflaherty/Documents/GRC/01-season-2026-27/GRC J16 2026-27 - Boys Season Intro.pptx")
print("boys deck saved")
